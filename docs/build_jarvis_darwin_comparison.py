from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUTPUT = "Jarvis_and_Darwin_Technical_Assessment.docx"

def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr(); el = OxmlElement("w:shd"); el.set(qn("w:fill"), fill); pr.append(el)

def border(cell):
    pr = cell._tc.get_or_add_tcPr(); bs = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}"); el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "6"); el.set(qn("w:color"), "D9D9D9"); bs.append(el)
    pr.append(bs)

def margins(cell, v=100, h=120):
    pr = cell._tc.get_or_add_tcPr(); mar = OxmlElement("w:tcMar")
    for name, val in (("top",v),("start",h),("bottom",v),("end",h)):
        el=OxmlElement(f"w:{name}"); el.set(qn("w:w"),str(val)); el.set(qn("w:type"),"dxa"); mar.append(el)
    pr.append(mar)

def body(doc, text, lead=None):
    p=doc.add_paragraph(style="Body Text")
    if lead and text.startswith(lead): p.add_run(lead).bold=True; p.add_run(text[len(lead):])
    else: p.add_run(text)
    return p

def bullets(doc, items):
    for item in items: doc.add_paragraph(item, style="List Bullet")

def table(doc, headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    trpr=t.rows[0]._tr.get_or_add_trPr(); trpr.append(OxmlElement("w:tblHeader"))
    for j,h in enumerate(headers):
        c=t.rows[0].cells[j]; c.text=h; shade(c,"263746"); c.width=Inches(widths[j]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(9)
    for i,row in enumerate(rows):
        cells=t.add_row().cells
        for j,val in enumerate(row):
            c=cells[j]; c.text=str(val); c.width=Inches(widths[j]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i%2: shade(c,"F2F6F8")
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.05
                for r in p.runs: r.font.size=Pt(8.5)
    for row in t.rows:
        for c in row.cells: border(c); margins(c)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11); sec.top_margin=Inches(.72); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(.82); sec.right_margin=Inches(.82)
styles=doc.styles
for n in ("Normal","Body Text","List Bullet"):
    styles[n].font.name="Aptos"; styles[n].font.size=Pt(10.5)
styles["Body Text"].paragraph_format.space_after=Pt(6); styles["Body Text"].paragraph_format.line_spacing=1.12
styles["List Bullet"].paragraph_format.space_after=Pt(3)
styles["Title"].font.name="Aptos Display"; styles["Title"].font.size=Pt(29); styles["Title"].font.bold=True; styles["Title"].font.color.rgb=RGBColor(0,0,0)
for n,s in (("Heading 1",18),("Heading 2",13),("Heading 3",11)):
    st=styles[n]; st.font.name="Aptos Display"; st.font.size=Pt(s); st.font.bold=True; st.font.color.rgb=RGBColor(0,0,0); st.paragraph_format.space_before=Pt(12); st.paragraph_format.space_after=Pt(5); st.paragraph_format.keep_with_next=True
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT; footer.add_run("Static technical assessment   |   13 September 2026").font.size=Pt(8)

p=doc.add_paragraph(style="Title"); p.add_run("Jarvis and Darwin Technical Assessment"); p.paragraph_format.space_before=Pt(65); p.paragraph_format.space_after=Pt(18)
p=doc.add_paragraph(); p.add_run("Architecture RAG voice stack security and engineering readiness").font.size=Pt(15); p.paragraph_format.space_after=Pt(26)
body(doc,"Jarvis is not better than Darwin as an engineering foundation for the voice problem Darwin is trying to solve. Jarvis is the stronger visual demonstration today: it has a polished browser HUD, broad Claude Code and MCP tool access, camera and gesture features, and sensible WebSocket streaming. Darwin has the more serious systems direction: a measurable audio front end, explicit realtime state, wake model evaluation, and a path toward reliable macOS operation.")
body(doc,"Jarvis does not contain retrieval augmented generation. It has neither an ingestion pipeline nor embeddings, a vector database, document chunking, retrieval ranking, citations, or a persistent knowledge index. Claude and MCP search tools may retrieve information during a turn, but that is tool use, not an application-owned RAG system. Conversation context also disappears when its WebSocket session is lost.")

doc.add_heading("Executive verdict",1)
table(doc,["Question","Finding","Confidence"],[
    ["Is Jarvis better than Darwin?","Better as a visual demo and integration shell; worse for Darwin’s core reliability objective.","High, within the supplied Darwin design and inspected Jarvis source"],
    ["Does Jarvis have RAG?","No. No RAG components were found in source or dependencies.","High"],
    ["Should Darwin adopt the Jarvis stack?","No wholesale rewrite. Borrow selected UI and bridge patterns only.","High"],
    ["Was Jarvis runtime-tested?","No. At the user’s direction, no dependencies were installed. Static syntax checks covered nine MJS files.","Exact"],
], [1.35,4.25,1.1])
body(doc,"Recommended decision. Keep Darwin’s Python and native macOS audio direction. If the web stack is disliked, do not inherit React, Three.js, Zustand, Framer Motion, MediaPipe, Porcupine Web, Kokoro, and the Node bridge as one package. Extract the useful ideas behind stable interfaces: a local agent bridge, streamed events, default-deny tool execution, a clean display contract, and explicit interruption behavior.","Recommended decision.")

doc.add_heading("Scope and method",1)
body(doc,"This assessment reviewed the public repository at commit 1c4016afdf86f7043efc6882ceffef84ad0d8783, dated 5 August 2026, together with the existing DARWIN Voice Agent Architecture and Implementation Plan in this workspace. The Jarvis checkout contained 64 tracked files and approximately 14,615 lines across 44 TypeScript, TSX, and MJS source files. The checkout was shallow, so repository maturity was not inferred from the one locally visible commit.")
bullets(doc,[
    "Inspected package manifest, bridge, voice and VAD code, browser client, state management, UI sanitization, media proxying, configuration, setup scripts, and documentation.",
    "Searched source and dependencies for retrieval, embeddings, vector storage, indexing, memory, and session persistence.",
    "Ran Node syntax checks on nine MJS files without loading project dependencies. All nine passed.",
    "Found no test or specification files and no test script. Build, lint, microphone behavior, browser permissions, Claude authentication, MCP actions, and end-to-end latency were not executed.",
])

doc.add_heading("How Jarvis is made",1)
body(doc,"Jarvis is two cooperating processes. A React and Vite browser application owns the visual interface, microphone interaction, local energy VAD, wake behavior, speech recognition selection, speech synthesis, camera capture, and gesture tracking. A Node bridge listens locally over WebSocket and HTTP, starts a Claude Agent SDK query, exposes local display and UI tools, imports MCP servers from the user’s Claude Code configuration, proxies selected media, and applies a coarse permission gate.")
table(doc,["Layer","Implementation","Assessment"],[
    ["Interface","React 19, Three.js through React Three Fiber, GLSL, Framer Motion, Zustand","Visually ambitious and demo-friendly; unusually large surface for a voice assistant."],
    ["Voice input","Browser SpeechRecognition fallback or ElevenLabs Scribe; optional Porcupine wake word; custom RMS energy VAD","Accessible setup, but browser and cloud behavior varies and the energy VAD is not an acoustic front end."],
    ["Voice output","Browser speechSynthesis, optional ElevenLabs, optional in-browser Kokoro ONNX","Good fallback concept; three engines increase behavioral and diagnostic complexity."],
    ["Agent runtime","Claude Agent SDK launched by a Node MJS bridge; one query session per WebSocket","Powerful reuse of Claude Code and MCP; tightly coupled to Anthropic tooling and socket lifetime."],
    ["Tool integration","Configured Claude Code MCP servers plus custom UI, browser, display, and camera servers","Broad capability with little bespoke integration work."],
    ["Transport","WebSocket for streaming conversation and UI events; HTTP for health, speech, file and media endpoints","Clear local boundary, but no durable session recovery."],
], [1.2,2.55,2.95])

doc.add_heading("RAG and memory assessment",1)
body(doc,"A real RAG implementation needs an owned knowledge corpus, ingestion and normalization, chunking, an embedding or sparse index, retrieval and ranking, context assembly, provenance, and an evaluation path. Jarvis implements none of these. Its model can call web search or other MCP tools, and Claude Code can read files available through tools, but the application does not turn those sources into a persistent retrievable knowledge base.")
table(doc,["Capability","Jarvis status","Evidence"],[
    ["Document ingestion and chunking","Absent","No ingestion service, chunk schema, or parser-to-index pipeline."],
    ["Embeddings and vector storage","Absent","No embedding SDK or vector database dependency; no corresponding source usage."],
    ["Retriever or reranker","Absent","No retrieval query path, scoring, filters, hybrid search, or reranking."],
    ["Provenance and citations","Absent as an application subsystem","Tools may return source material, but Jarvis has no citation record or retrieval trace."],
    ["Persistent conversation memory","Absent","Source states that the socket is the session and reconnecting creates a fresh session."],
    ["Transient model context","Present","The live Claude Agent SDK query carries the current conversation until the socket closes."],
], [1.55,1.55,3.6])
body(doc,"If Darwin needs RAG, it should be a separate service behind the conversation runtime, not embedded into audio code. Start only with a defined corpus and use case. Add source identifiers, access controls, chunk versioning, hybrid retrieval, citations, and a small question set that measures retrieval recall and answer groundedness. RAG will not fix wake-word or barge-in failures.")

doc.add_heading("Direct comparison",1)
table(doc,["Criterion","Jarvis","Darwin","Advantage"],[
    ["Primary goal","Cinematic multimodal assistant demo","Reliable macOS voice agent under real acoustic conditions","Different goals"],
    ["Wake path","Browser SpeechRecognition or Porcupine Web","Owned ONNX wake model with condition-specific evaluation plan","Darwin for control and measurement"],
    ["Acoustic front end","Custom energy VAD; no implemented AEC, NS, AGC, render reference, or synchronized audio telemetry","Native AVAudioEngine or AUVoiceIO AEC, NS, AGC and replay harness proposed","Darwin architecture"],
    ["Barge-in","Energy onset stops speech quickly","Planned cleaned-audio interruption confirmation with exact played-audio ledger","Jarvis is simpler today; Darwin is safer if implemented"],
    ["Agent and tools","Claude Code plus broad MCP reuse","Modular conversation runtime and existing integrations","Jarvis for breadth today"],
    ["UI","Rich WebGL HUD, panels, camera, gestures","Not the current architectural focus","Jarvis"],
    ["RAG","None","Not described as implemented in supplied plan","Neither"],
    ["Testing","No automated tests found","Explicit deterministic wake, soak, latency, and route test plan","Darwin direction"],
    ["Persistence","Conversation lost on socket drop","Structured local event log and replay recommended","Darwin direction"],
    ["Security","Origin checks, SSRF guards, HTML sanitization, coarse preconfigured write gate","Strong confirmation or OS authentication recommended for sensitive actions","Darwin policy; Jarvis has useful implementation pieces"],
], [1.05,2.1,2.75,.8])

doc.add_heading("Engineering quality and risk",1)
doc.add_heading("What Jarvis does well",2)
bullets(doc,[
    "The source is unusually candid about browser failure modes, reconnect behavior, audio fallbacks, latency, and security boundaries.",
    "Local-origin checks, SSRF target vetting, strict content security policy, DOM sanitization, media host restrictions, and real-path confinement show thoughtful defensive work.",
    "The permission default is read-oriented, while write-capable built-ins and effectful MCP names require an explicit process setting.",
    "Streaming UI events, sentence-level speech, interruption calls, diagnostics, and capability probing are good interaction patterns.",
])
doc.add_heading("What prevents a production recommendation",2)
bullets(doc,[
    "There are no automated tests, test runner, coverage configuration, CI workflow, or reproducible acceptance results in the inspected repository.",
    "The permission system classifies MCP tools partly by server names and verb regular expressions. Tool names are not a security contract; metadata-driven capabilities and per-action confirmation are safer.",
    "Enabling writes is process-wide. Voice commands that send, buy, delete, publish, or operate signed-in Chrome need transaction-specific confirmation and often operating-system authentication.",
    "The browser microphone path has no acoustic echo cancellation boundary owned and measured by the application. Energy-only VAD can fire on music, impact sounds, assistant playback, and environmental changes.",
    "Conversation durability is knowingly absent. The visible transcript can survive while the model forgets it, creating a misleading state split.",
    "Direct browser mode embeds provider and integration credentials into the client bundle; the code warns that it is unsuitable for public deployment.",
    "Configuration documentation and source disagree on the default bridge effort: the README says medium while server.mjs uses high. That signals release drift.",
    "The large frontend dependency set increases update, bundle, browser compatibility, GPU, permission, and supply-chain burden for capabilities unrelated to core voice reliability.",
])

doc.add_heading("Tech stack recommendation",1)
body(doc,"Your dislike of the stack is justified if the objective is a dependable personal macOS assistant rather than a cinematic browser demo. Jarvis combines several frameworks because its interface is the product. Darwin’s product risk is elsewhere: capture quality, device routing, wake recall, endpointing, cancellation, and safe actions. A rewrite into the Jarvis stack would move effort away from those risks.")
table(doc,["Keep or borrow","Avoid importing","Why"],[
    ["Local agent bridge contract","React and Three.js as a requirement","The bridge is reusable; the rendering framework is not."],
    ["Typed streamed events","Socket-bound memory","Events support observability, while memory needs durable IDs and replay."],
    ["Capability probing and diagnostics","Three overlapping TTS paths at launch","Fallbacks are useful after the primary path is measured."],
    ["Origin checks, SSRF protection, sanitization","Regex-only action authorization","Security controls are valuable; authorization needs declared tool effects and confirmations."],
    ["MCP adapter layer","Anthropic-specific logic throughout the application","Keep providers behind one conversation interface."],
], [2.0,2.05,2.65])
body(doc,"Preferred Darwin shape. Use Swift for the native macOS audio and permission boundary; keep the current Python services where they already work for wake inference, evaluation, and orchestration; use local IPC with versioned events; add a thin provider adapter for Claude, OpenAI, or local models; and choose a native SwiftUI interface or a minimal web view only if the UI genuinely needs web rendering.","Preferred Darwin shape.")

doc.add_heading("Recommended adoption plan",1)
table(doc,["Priority","Action","Outcome"],[
    ["1","Continue Darwin’s native voice-processing spike and deterministic audio evaluation harness.","Resolves the highest-risk acoustic failure with evidence."],
    ["2","Define a versioned event protocol for speech start, final transcript, tool proposal, response audio, interruption, and completion.","Lets audio, agent, UI, and telemetry evolve independently."],
    ["3","Prototype a Claude Agent SDK or equivalent adapter behind the conversation interface, borrowing Jarvis’s streaming bridge pattern.","Gains MCP breadth without committing Darwin to the full frontend stack."],
    ["4","Replace a global write switch with tool-effect metadata, scoped grants, visible confirmation, and OS authentication for sensitive actions.","Makes voice-driven automation defensible."],
    ["5","Add durable session IDs and append-only turn events, including what speech was actually played.","Supports crash recovery, replay, evaluation, and consistent memory."],
    ["6","Introduce RAG only after selecting a bounded corpus and evaluation set.","Avoids building an expensive feature with no measurable retrieval job."],
    ["7","Borrow the HUD concept selectively after reliability gates pass.","Keeps visual polish from dictating the system architecture."],
], [.55,4.35,1.8])

doc.add_heading("Final answer",1)
body(doc,"Jarvis is an impressive interface and Claude Code integration demo, not a stronger voice-system architecture than Darwin. It is worth studying for its streamed local bridge, MCP reuse, UI tool protocol, browser-origin defense, media sanitization, and candid fallback handling. It should not replace Darwin’s native audio plan, evaluation discipline, or modular architecture.")
body(doc,"There is no RAG in Jarvis. There is transient Claude conversation context and runtime tool retrieval, but no owned retrieval system and no durable memory. The best course is selective adoption: take the bridge and interaction patterns, keep Darwin’s audio and evaluation core, simplify the UI stack, strengthen action authorization, and defer RAG until it has a precise knowledge task and measurable acceptance tests.")

doc.add_heading("Sources and evidence",1)
bullets(doc,[
    "Jarvis repository: https://github.com/adewaskar/jarvis",
    "Inspected Jarvis commit: https://github.com/adewaskar/jarvis/commit/1c4016afdf86f7043efc6882ceffef84ad0d8783",
    "Local source checkout used for static review: /private/tmp/jarvis-review",
    "DARWIN Voice Agent Architecture and Implementation Plan, local workspace document, September 2026.",
])
body(doc,"Evidence limitation. This is a source-level assessment, not a runtime certification. At the user’s request no dependencies were installed. Claims about successful microphone capture, speech providers, Claude login, MCP tools, Chrome control, GPU rendering, latency, and end-to-end behavior remain unverified until the application is run in its supported browser and environment.","Evidence limitation.")

doc.save(OUTPUT)
print(OUTPUT)
